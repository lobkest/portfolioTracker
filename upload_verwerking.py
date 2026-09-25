"""
Upload-verwerking: de taakfuncties achter POST /upload -- Excel inlezen,
kolom-normalisatie, ticker-resolutie (twee paden: 'niet opslaan' en
opslaand), Order ID-bepaling, portfolio-code-matching, DB-insert en
dividend-bestand-verwerking. _upload_impl() in app.py orkestreert deze
taakfuncties in de juiste volgorde.

Losgetrokken uit app.py.
"""
import hashlib

import pandas as pd
import openpyxl
from flask import request

from debug_utils import dprint, meet_tijd
from ticker_zekerheid import (
    basis_ticker_zekerheid_parallel, vind_tickers_met_snelle_prijscheck_parallel,
    find_ticker_met_snelle_prijscheck,
)
from portfolio_admin import find_matching_code, generate_code
from db import save_dividenden
from dividend import verwerk_rekeningoverzicht

# Kolomnaam exact zoals DeGiro 'm in het transactiebestand zet (na
# df.columns.str.strip(), dat evt. rondom-spaties in de header wegwerkt).
# Ontbreekt in oudere DeGiro-exportformaten — daarom overal met een
# beschikbaarheids-check behandeld i.p.v. als verplichte kolom.
KOSTEN_KOLOM = "Transactiekosten en/of kosten van derden EUR"

# Kale waarde (aantal x koers, zonder AutoFX/transactiekosten) — DEGIRO's
# eigen GAK-weergave is hierop gebaseerd, in tegenstelling tot Totaal EUR
# (dat wel kosten meetelt en de GAK structureel te hoog maakt, zie
# CLAUDE.md). Zelfde beschikbaarheids-check-patroon als KOSTEN_KOLOM.
WAARDE_KOLOM = "Waarde EUR"

# DEGIRO's eigen afrekenkoers voor deze transactie — preciezer dan een losse
# historische FX-lookup achteraf. Gebruikt om de rauwe 'Koers'-kolom (die
# voor een niet-EUR-genoteerde positie, bv. TTWO op NDQ, gewoon de
# vreemde-valuta-koers bevat) naar EUR om te rekenen vóór opslag — zie
# CLAUDE.md, Databasestructuur (transacties.koers). Leeg/NaN voor
# EUR-genoteerde rijen. Zelfde beschikbaarheids-check-patroon als
# KOSTEN_KOLOM/WAARDE_KOLOM.
WISSELKOERS_KOLOM = "Wisselkoers"


def _normaliseer_tijd(waarde):
    """Zet de 'Tijd'-kolom uit het transactiebestand om naar een string die
    Postgres' TIME-kolom kan opslaan. Pandas/openpyxl kan een tijdcel als
    string ("13:39"), datetime.time of datetime.datetime teruggeven,
    afhankelijk van hoe de cel in Excel geformatteerd is."""
    if pd.isna(waarde):
        return None
    if hasattr(waarde, "strftime"):
        return waarde.strftime("%H:%M:%S")
    return str(waarde)

def _lees_transacties_excel(bestand1):
    """Taak 1/6 (Excel inlezen): leest het transactiebestand in en zet de
    kolomnamen/Datum-kolom in bruikbare vorm."""
    bestand1.seek(0)
    df = pd.read_excel(bestand1)
    df.columns = df.columns.str.strip()
    df["Datum"] = pd.to_datetime(df["Datum"], dayfirst=True)
    return df


def _normaliseer_transactie_kolommen(df):
    """Taak 2/6 (kolom-normalisatie): voegt _kosten_eur/_waarde_eur/_koers_eur
    toe op basis van KOSTEN_KOLOM/WAARDE_KOLOM/WISSELKOERS_KOLOM -- met een
    nette fallback (NaN resp. ongewijzigde Koers-kolom) als die kolommen in
    dit DeGiro-exportformaat ontbreken."""
    if KOSTEN_KOLOM in df.columns:
        df["_kosten_eur"] = pd.to_numeric(df[KOSTEN_KOLOM], errors="coerce")
    else:
        df["_kosten_eur"] = pd.Series([None] * len(df), index=df.index, dtype="float64")

    if WAARDE_KOLOM in df.columns:
        df["_waarde_eur"] = pd.to_numeric(df[WAARDE_KOLOM], errors="coerce")
    else:
        df["_waarde_eur"] = pd.Series([None] * len(df), index=df.index, dtype="float64")

    if WISSELKOERS_KOLOM in df.columns:
        wisselkoers = pd.to_numeric(df[WISSELKOERS_KOLOM], errors="coerce")
        heeft_wisselkoers = wisselkoers.notna() & (wisselkoers != 0)
        df["_koers_eur"] = df["Koers"].astype(float)
        df.loc[heeft_wisselkoers, "_koers_eur"] = (
            df.loc[heeft_wisselkoers, "Koers"].astype(float) / wisselkoers.loc[heeft_wisselkoers]
        )
        for _, rij in df.loc[heeft_wisselkoers, ["ISIN", "Beurs", "Koers", WISSELKOERS_KOLOM, "_koers_eur"]].iterrows():
            dprint(
                f"[koers-eur] ISIN={rij['ISIN']} Beurs={rij['Beurs']}: "
                f"Koers={rij['Koers']} / Wisselkoers={rij[WISSELKOERS_KOLOM]} -> "
                f"_koers_eur={rij['_koers_eur']:.4f}"
            )
    else:
        df["_koers_eur"] = df["Koers"].astype(float)
        dprint(f"[upload] WAARSCHUWING: kolom '{WISSELKOERS_KOLOM}' niet gevonden - "
               f"koers-kolom blijft ongewijzigd (aanname: al EUR)")
    return df


def _ticker_resolutie_niet_opslaan_pad(df):
    """Taak 3/6 (ticker-resolutie 'niet opslaan'-pad): lichte, parallelle
    ticker-zekerheid per (ISIN, Beurs)-groep -- zie de uitgebreide toelichting
    in _upload_impl() voor waarom dit bewust de goedkope variant is (geen
    volledige verifieer_tickers_met_prijs_parallel(), zie CLAUDE.md,
    Statistieken-incident 2026-08-31) """
    groepen = list(df.groupby(["ISIN", "Beurs"]))
    namen = [groep["Product"].iloc[0] for (_isin, _beurs_val), groep in groepen]

    with meet_tijd(f"ticker_resolutie_niet_opslaan ({len(groepen)} positie(s))"):
        posities_voor_check = [
            (naam_positie, isin, beurs_val, [
                {"datum": row["Datum"].strftime("%Y-%m-%d"), "koers": float(row["_koers_eur"])}
                for _, row in groep.iterrows()
            ])
            for naam_positie, ((isin, beurs_val), groep) in zip(namen, groepen)
        ]
        resultaten = basis_ticker_zekerheid_parallel(posities_voor_check)

        ticker_by_isin_beurs = {}
        ticker_zekerheid = []
        ticker_posities_ruw = []
        for (naam_positie, isin, beurs_val, transacties_lijst), resultaat in zip(posities_voor_check, resultaten):
            resultaat["isin"] = isin
            resultaat["naam"] = naam_positie
            resultaat["echte_naam"] = naam_positie
            ticker_by_isin_beurs[(isin, beurs_val)] = resultaat["ticker"]
            ticker_zekerheid.append(resultaat)
            ticker_posities_ruw.append({
                "naam": naam_positie, "isin": isin, "beurs": beurs_val,
                "transacties": transacties_lijst,
            })

    return ticker_by_isin_beurs, ticker_zekerheid, ticker_posities_ruw


def _bouw_transacties_df_niet_opslaan(df, ticker_by_isin_beurs):
    """Bouwt de transacties_df voor het 'niet opslaan'-pad uit de ruwe
    Excel-df + de opgeloste tickers -- zelfde kolomvorm als wat normaal uit
    de database komt, zodat analyze_transacties() ongewijzigd kan blijven."""
    return pd.DataFrame({
        "datum": df["Datum"],
        "product": df["Product"],
        "isin": df["ISIN"],
        "beurs": df["Beurs"],
        "ticker": [ticker_by_isin_beurs.get((isin_val, beurs_val))
                   for isin_val, beurs_val in zip(df["ISIN"], df["Beurs"])],
        "aantal": df["Aantal"].astype(float),
        "koers": df["_koers_eur"].astype(float),
        "totaal_eur": df["Totaal EUR"].astype(float),
        "echte_naam": df["Product"],
        "transactiekosten": df["_kosten_eur"],
        "waarde_eur": df["_waarde_eur"],
        "tijd": df["Tijd"],
    })


def _bepaal_order_ids(bestand1, df):
    """Order ID-kolom kan door merged cells één kolom verschoven staan
    t.o.v. de header; leest 'm daarom apart uit met openpyxl, die de
    waarden onder de merge vindt. Rijen zonder echte (UUID-vormige) Order
    ID krijgen daarna een synthetische, stabiele ID."""
    bestand1.seek(0)
    wb = openpyxl.load_workbook(bestand1, data_only=True)
    ws = wb.active
    order_ids_ruw = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        gevonden = None
        for cell in row:
            if cell.value and isinstance(cell.value, str) and len(cell.value) == 36 and cell.value.count("-") == 4:
                gevonden = cell.value
                break
        order_ids_ruw.append(gevonden)

    if len(order_ids_ruw) == len(df):
        df["Order ID"] = order_ids_ruw
    else:
        df["Order ID"] = None

    def basis_hash(row):
        basis = f"{row['Datum']}|{row['Tijd']}|{row['Product']}|{row['ISIN']}|{row['Aantal']}|{row['Totaal EUR']}"
        return "SYN-" + hashlib.md5(basis.encode()).hexdigest()[:16]

    heeft_order_id = df["Order ID"].notna()
    if (~heeft_order_id).any():
        synthetische_ids = df.loc[~heeft_order_id].apply(basis_hash, axis=1)
        volgnummer = synthetische_ids.groupby(synthetische_ids).cumcount()
        synthetische_ids = synthetische_ids + "-" + volgnummer.astype(str)
        df.loc[~heeft_order_id, "Order ID"] = synthetische_ids

    return df


def _vind_of_maak_portfolio_code(cur, df, naam):
    """Zoekt een bestaande portfolio die dezelfde persoon vertegenwoordigt
    (via Order ID-overlap, zie find_matching_code) of maakt een nieuwe code
    aan. Geeft (code, match_code, rows_to_insert) terug."""
    new_order_ids = set(df["Order ID"])
    match_code, missing_ids = find_matching_code(cur, new_order_ids)

    if match_code:
        code = match_code
        if naam:
            cur.execute("UPDATE portfolios SET naam = %s WHERE code = %s", (naam, code))
        rows_to_insert = df[df["Order ID"].isin(missing_ids)] if missing_ids else df.iloc[0:0]
    else:
        code = generate_code(cur)
        cur.execute("INSERT INTO portfolios (code, naam) VALUES (%s, %s)", (code, naam or None))
        rows_to_insert = df

    return code, match_code, rows_to_insert


def _ticker_resolutie_opslaan_pad(cur, code, rows_to_insert, herbepaal_alle_tickers):
    """Taak 4/6 (ticker-resolutie opslaan-pad): lichte, parallelle
    prijscontrole per (ISIN, Beurs)-groep, met hergebruik van al bekende
    tickers tenzij het 'ticker-informatie opnieuw bepalen'-vinkje aan
    staat (zie CLAUDE.md). Geeft ticker_by_isin_beurs terug."""
    groepen = list(rows_to_insert.groupby(["ISIN", "Beurs"]))

    bekende_tickers = {}
    if not herbepaal_alle_tickers:
        cur.execute(
            "SELECT isin, beurs, ticker FROM transacties WHERE code = %s AND ticker IS NOT NULL",
            (code,),
        )
        bekende_tickers = {(isin_val, beurs_val): ticker for isin_val, beurs_val, ticker in cur.fetchall()}

    transacties_per_groep = {
        key: [
            {"datum": row["Datum"].strftime("%Y-%m-%d"), "koers": float(row["_koers_eur"])}
            for _, row in groep.iterrows()
        ]
        for key, groep in groepen
    }
    eerste_poging = [
        (groep["Product"].iloc[0], key[0], key[1], transacties_per_groep[key])
        for key, groep in groepen
    ]
    resultaten = vind_tickers_met_snelle_prijscheck_parallel(eerste_poging, bekende_tickers=bekende_tickers)

    ticker_by_isin_beurs = {}
    for (key, groep), detail in zip(groepen, resultaten):
        if not detail["ticker"]:
            # Zeldzame fallback: de eerste Product-naam van de groep gaf
            # geen match, probeer de overige rijen (zelfde gedrag als
            # voorheen). Goedkoop: zonder ticker doet find_ticker_met_
            # snelle_prijscheck() geen enkele prijscheck.
            for _, row in groep.iterrows():
                detail = find_ticker_met_snelle_prijscheck(
                    row["Product"], row["ISIN"], row["Beurs"], transacties_per_groep[key],
                    bekende_tickers.get(key),
                )
                if detail["ticker"]:
                    break
        ticker_by_isin_beurs[key] = detail["ticker"]

    return ticker_by_isin_beurs


def _insert_nieuwe_transacties(cur, code, rows_to_insert, ticker_by_isin_beurs):
    """Taak 5/6 (DB-insert): voegt nieuwe transactierijen in;
    ON CONFLICT DO NOTHING negeert rijen die (op order_id) al bestaan.
    Geeft het aantal INSERT-pogingen zonder exception terug (dus inclusief
    rijen die door ON CONFLICT genegeerd werden)."""
    ingevoegd = 0
    for _, row in rows_to_insert.iterrows():
        try:
            kosten_waarde = row["_kosten_eur"]
            waarde_eur_waarde = row["_waarde_eur"]
            cur.execute(
                """INSERT INTO transacties
                   (code, datum, product, isin, beurs, ticker, aantal, koers, totaal_eur, order_id, echte_naam, transactiekosten, waarde_eur, tijd)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                   ON CONFLICT (code, order_id) DO NOTHING""",
                (code, row["Datum"].date(), row["Product"], row["ISIN"], row["Beurs"],
                 ticker_by_isin_beurs[(row["ISIN"], row["Beurs"])], float(row["Aantal"]), float(row["_koers_eur"]),
                 float(row["Totaal EUR"]), row["Order ID"], row["Product"],
                 float(kosten_waarde) if pd.notna(kosten_waarde) else None,
                 float(waarde_eur_waarde) if pd.notna(waarde_eur_waarde) else None,
                 _normaliseer_tijd(row["Tijd"])),
            )
            ingevoegd += 1
        except Exception as e:
            pass
    return ingevoegd


def _verwerk_dividend_bestand_indien_aanwezig(code):
    """Taak 6/6 (dividend-verwerking): leest het optionele Account-
    rekeningoverzicht-bestand (request.files['bestand2']) en slaat de
    dividenden op. Doet niets als er geen bestand2 is meegestuurd."""
    bestand2 = request.files.get("bestand2")
    if bestand2 and bestand2.filename != "":
        with meet_tijd("dividend_bestand_verwerken"):
            dividend_records = verwerk_rekeningoverzicht(bestand2)
            save_dividenden(code, dividend_records)
